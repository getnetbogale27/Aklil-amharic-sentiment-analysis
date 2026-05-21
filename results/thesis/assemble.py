"""Main assembler — runs all chapter builders and saves the docx."""

import sys, os, json
from datetime import datetime

sys.path.insert(0, "/home/user/Aklil-amharic-sentiment-analysis/results/thesis")

from docx import Document
from build_thesis import (set_page_size, build_title_page, build_declaration,
                           build_certification, build_acknowledgment, build_abstract,
                           build_abbreviations, build_toc, build_list_of_tables,
                           build_list_of_figures, OUT, FIG)
from build_chapters123 import build_chapter1
from build_ch2 import build_chapter2
from build_ch3 import build_chapter3
from build_ch4 import build_chapter4
from build_ch5_refs_app import build_chapter5, build_references, build_appendices

def count_words(doc):
    total = 0
    for para in doc.paragraphs:
        total += len(para.text.split())
    return total

def main():
    print("Building thesis document...")
    doc = Document()
    set_page_size(doc)

    # Remove default empty paragraph
    for para in doc.paragraphs:
        p = para._p
        p.getparent().remove(p)
        break

    print("  Building front matter...")
    build_title_page(doc)
    build_declaration(doc)
    build_certification(doc)
    build_acknowledgment(doc)
    build_abstract(doc)
    build_abbreviations(doc)
    build_toc(doc)
    build_list_of_tables(doc)
    build_list_of_figures(doc)

    print("  Building Chapter 1...")
    build_chapter1(doc)

    print("  Building Chapter 2...")
    build_chapter2(doc)

    print("  Building Chapter 3...")
    build_chapter3(doc)

    print("  Building Chapter 4...")
    build_chapter4(doc)

    print("  Building Chapter 5...")
    build_chapter5(doc)

    print("  Building References...")
    build_references(doc)

    print("  Building Appendices...")
    build_appendices(doc)

    # Save
    os.makedirs(OUT, exist_ok=True)
    today = datetime.now().strftime("%Y%m%d")
    filename = f"Aklilu_Gebeyehu_MSc_Thesis_{today}.docx"
    out_path = os.path.join(OUT, filename)
    doc.save(out_path)
    print(f"  Saved: {out_path}")

    # Build log
    word_count = count_words(doc)
    para_count = len(doc.paragraphs)
    estimated_pages = max(1, para_count // 4)

    # Count figures found vs missing
    figs = [
        "fig3_1_methodology_pipeline.jpeg",
        "fig4_1_label_distribution.jpeg",
        "fig4_2_tweet_length_by_class.jpeg",
        "fig4_3_char_vs_word_ngrams.jpeg",
        "fig4_4_all_models_comparison.jpeg",
        "fig4_5_best_model_cm.jpeg",
        "fig4_6_per_class_f1_heatmap.jpeg",
        "fig4_7_experiment_progression.jpeg",
        "fig4_8_bilstm_training_curves.jpeg",
        "fig4_9_learning_curve.jpeg",
    ]
    found = [f for f in figs if os.path.exists(os.path.join(FIG, f))]
    missing = [f for f in figs if not os.path.exists(os.path.join(FIG, f))]

    build_log = {
        "timestamp": datetime.now().isoformat(),
        "output_file": out_path,
        "total_word_count": word_count,
        "total_paragraphs": para_count,
        "estimated_pages": estimated_pages,
        "figures_found": found,
        "figures_missing": missing,
        "figures_inserted": len(found),
        "tables_inserted": sum(1 for t in doc.tables),
        "chapters": {
            "front_matter": ["Title Page", "Declaration", "Certification",
                             "Acknowledgment", "Abstract", "Abbreviations",
                             "TOC", "List of Tables", "List of Figures"],
            "chapter1": "Introduction",
            "chapter2": "Review of Related Literature",
            "chapter3": "Research Methodology",
            "chapter4": "Results and Discussion",
            "chapter5": "Conclusion and Future Work",
            "references": "References",
            "appendices": ["A", "B", "C", "D", "E"],
        }
    }

    log_path = os.path.join(OUT, "thesis_build_log.json")
    with open(log_path, "w") as f:
        json.dump(build_log, f, indent=2)
    print(f"  Build log: {log_path}")
    print(f"\nDone! Word count: {word_count:,}, Est. pages: {estimated_pages}")
    print(f"  Figures inserted: {len(found)}/{len(figs)}")
    print(f"  Tables inserted: {sum(1 for t in doc.tables)}")
    if missing:
        print(f"  Missing figures: {missing}")

if __name__ == "__main__":
    main()
